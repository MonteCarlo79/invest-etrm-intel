"""系统运行费 extraction from 代购价格 sources (backfill + article automation).

Two source shapes, one writer:
- extract_from_image_bytes(img_bytes) -> float|None : Claude vision reads the
  per-province 代理购电价格表 and returns 系统运行费用折合度电水平 (元/kWh).
- extract_from_pdf_bytes(pdf_bytes) -> float|None : text-layer first (the row is
  plain text in official PDFs), vision fallback via pdfplumber page render.
- upsert_sysopfee(province, year_month, fee_yuan_kwh, source_file) : idempotent
  upsert into province_sysopfee_monthly (the 系统运行费 tab's table).

Runs on ECS (hermes) — LLM calls are ECS-only by design (Anthropic geo-block
from China Mac). The vision model is claude-sonnet-4-6 via shared.anthropic_client.
"""
from __future__ import annotations

import base64
import logging
import re
from typing import Optional

logger = logging.getLogger(__name__)

_PROMPT = (
    "这是一张中国省级电网代理购电价格公告的表格图片。请找到价格明细表中"
    "「系统运行费用折合度电水平」这一行（也可能写作「系统运行费折合度电水平」"
    "或「系统运行费」），返回其数值。\n"
    "注意单位：有的省份用「元/千瓦时」，有的用「分/千瓦时」。"
    "如果是分/千瓦时，请除以100换算成元/千瓦时后再输出。\n"
    "只输出一个数字（如 0.0357），不要任何其他文字。找不到就输出 NONE。"
)

# Text-layer patterns for PDFs (no vision needed when text exists).
# NOTE: the table line carries a 计算关系 formula cell between the label and the
# value, e.g. "8 系统运行费用折合度电水平 8=9+10+11+12+13+14+15+16 0.044318".
# Capturing the first digit after the label grabs the formula index (the
# 2026-05 incident: 上海 read as 8.0 instead of 0.044318). So: take the text
# after the label, strip formula tokens, take the LAST number on the line.
_LABEL_RE = re.compile(r"系统运行费[用]?折合度电水平(.*)")
# Formula tokens look like "8=9+10+11" or "+12+13+14" — a bare number ("0" in
# "0.0357") is NOT a formula; require '=' or at least two consecutive +N terms.
_FORMULA_TOKEN_RE = re.compile(r"\d+(?:=\d+)+(?:\+\d+)*|(?:\+\d+){2,}")
_NUMBER_RE = re.compile(r"-?\d+\.\d+|-?\d+")


def _extract_fee_from_line(line: str) -> Optional[float]:
    """Extract the fee value from a text line containing the 折合度电水平 label.
    Returns None when the line carries no usable number (bare label header)."""
    m = _LABEL_RE.search(line)
    if not m:
        return None
    rest = _FORMULA_TOKEN_RE.sub(" ", m.group(1))
    nums = _NUMBER_RE.findall(rest)
    if not nums:
        return None
    return _normalize_fee(float(nums[-1]))


_MAX_IMAGE_BYTES = 4_500_000  # API hard limit is 5 MB; leave headroom


def _sniff_media_type(data: bytes) -> str:
    """File extensions lie (a '.jpg' may be WebP). Sniff magic bytes."""
    if data[:4] == b"\x89PNG":
        return "image/png"
    if data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return "image/webp"
    if data[:2] == b"\xff\xd8":
        return "image/jpeg"
    if data[:6] in (b"GIF87a", b"GIF89a"):
        return "image/gif"
    return "image/jpeg"


def _shrink_image(image_bytes: bytes) -> bytes:
    """Re-encode to fit under the vision API's 5 MB limit (high-res phone photos
    of price tables arrive at 6-10 MB). Returns input unchanged if small enough."""
    if len(image_bytes) <= _MAX_IMAGE_BYTES:
        return image_bytes
    import io

    from PIL import Image
    img = Image.open(io.BytesIO(image_bytes))
    if max(img.size) > 2400:
        img.thumbnail((2400, 2400))
    buf = io.BytesIO()
    img.convert("RGB").save(buf, format="JPEG", quality=85)
    return buf.getvalue()


def _vision_extract(image_bytes: bytes, media_type: str, api_key: str) -> Optional[float]:
    from shared.anthropic_client import make_client
    image_bytes = _shrink_image(image_bytes)
    media_type = _sniff_media_type(image_bytes)  # extensions lie; magic bytes don't
    client = make_client(api_key)
    resp = client.messages.create(
        model="claude-sonnet-4-6",
        max_tokens=64,
        messages=[{
            "role": "user",
            "content": [
                {"type": "image",
                 "source": {"type": "base64", "media_type": media_type,
                            "data": base64.standard_b64encode(image_bytes).decode()}},
                {"type": "text", "text": _PROMPT},
            ],
        }],
    )
    text = resp.content[0].text.strip()
    if text == "NONE":
        return None
    m = re.search(r"-?\d+\.?\d*", text)
    return _normalize_fee(float(m.group(0))) if m else None


def _normalize_fee(fee: float) -> float:
    """Guard against 分/千瓦时 tables: no province's 系统运行费 exceeds 1 元/kWh
    (observed range 0.02–0.15); a value > 1 means the table was in 分."""
    return fee / 100.0 if fee > 1.0 else fee


def extract_from_image_bytes(image_bytes: bytes, *, media_type: str = "image/jpeg",
                             api_key: str) -> Optional[float]:
    """Vision-extract 系统运行费用折合度电水平 (元/kWh) from a price-table image."""
    return _vision_extract(image_bytes, media_type, api_key)


def extract_from_pdf_bytes(pdf_bytes: bytes, *, api_key: str) -> Optional[float]:
    """Text-layer first (line-based, formula-cell aware); vision fallback."""
    import io
    import pdfplumber
    with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
        lines = [l for p in pdf.pages for l in (p.extract_text() or "").splitlines()]
    for line in lines:
        fee = _extract_fee_from_line(line)
        if fee is not None:
            return fee
    # vision fallback: render pages until a value is found (the price table is
    # not always page 1 — 浙江's is page 4)
    with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
        for page in pdf.pages[:6]:
            img = page.to_image(resolution=170).original
            buf = io.BytesIO()
            img.save(buf, format="PNG")
            fee = _vision_extract(buf.getvalue(), "image/png", api_key)
            if fee is not None:
                return fee
    return None


def upsert_sysopfee(cur, province: str, year_month: str, fee_yuan_kwh: float,
                    source_file: str) -> None:
    """Idempotent upsert into province_sysopfee_monthly (the 系统运行费 tab table)."""
    cur.execute(
        """
        INSERT INTO province_sysopfee_monthly
            (province, year_month, fee_yuan_kwh, source_file)
        VALUES (%s, %s, %s, %s)
        ON CONFLICT (province, year_month) DO UPDATE SET
            fee_yuan_kwh = EXCLUDED.fee_yuan_kwh,
            source_file = EXCLUDED.source_file,
            ingested_at = NOW()
        """,
        (province, year_month, fee_yuan_kwh, source_file),
    )


# ── Article-driven automation (daily 智慧光伏 scrape) ─────────────────────────
# Monthly 代理购电价格 announcement articles carry one price-table image per
# province. news_screener preserves <img> as [IMG:url] markers in the body
# (mmbiz.qpic.cn CDN — reachable from ECS without the WeChat proxy).
# For each image: vision-extract {province, fee} and upsert the article month.

_PROMPT_PROVINCE_FEE = (
    "这是一张中国省级电网代理购电价格公告的表格图片。请识别：\n"
    "1. 这是哪个省份/地区的公告（表头或标题通常包含省名）。省份用短名回答"
    "（如：山东、山西、蒙西、蒙东、冀北、深圳、新疆、甘肃；北京/上海/天津/重庆不要带「市」字）。\n"
    "2. 价格明细表中「系统运行费用折合度电水平」这一行（也可能写作"
    "「系统运行费折合度电水平」或「系统运行费」）的数值。\n"
    "注意单位：有的省份用「元/千瓦时」，有的用「分/千瓦时」。"
    "如果是分/千瓦时，请除以100换算成元/千瓦时后再输出。\n"
    "只输出 JSON：{\"province\": \"山东\", \"fee\": 0.0357}。"
    "找不到系统运行费数值就输出 {\"province\": \"...\", \"fee\": null}。不要输出任何其他文字。"
)

_IMG_MARKER_RE = re.compile(r"\[IMG:(https?://[^\]\s]+)\]")
_TITLE_MONTH_RE = re.compile(r"(20\d{2})\s*年\s*(\d{1,2})\s*月")
_ARTICLE_GATE_RE = re.compile(r"代理购电|代购电")
_QPIC_HOST_RE = re.compile(r"^https?://mmbiz\.q(pic|logo)\.cn/")

_MAX_IMAGES_PER_ARTICLE = 40
_MIN_IMAGE_BYTES = 15_000  # skip icons / dividers / QR codes

_STRIP_SUFFIXES = (
    "壮族自治区", "回族自治区", "维吾尔自治区", "特别行政区", "自治区", "省", "市",
)

# Canonical province set = what's already in province_sysopfee_monthly (DB probe
# 2026-10-31, 34 names). Anything outside this list (e.g. 闽南, a region of
# 福建 that scroll tables sometimes split out) is rejected.
_PROVINCE_CANON = {
    "北京", "天津", "上海", "重庆", "河北", "冀北", "河北南网", "山西", "山东",
    "江苏", "浙江", "安徽", "福建", "江西", "河南", "湖北", "湖南", "广东",
    "深圳", "广西", "海南", "贵州", "云南", "四川", "陕西", "甘肃", "青海",
    "宁夏", "新疆", "辽宁", "吉林", "黑龙江", "蒙西", "蒙东",
}


def _canon_province_name(name: str) -> Optional[str]:
    """Normalize a model-returned province name to the platform short name."""
    n = (name or "").strip()
    for suf in _STRIP_SUFFIXES:
        if n.endswith(suf):
            n = n[: -len(suf)]
            break
    if n in {"内蒙古", ""}:  # ambiguous — must be 蒙西/蒙东
        return None
    return n if n in _PROVINCE_CANON else None


# ── Scroll-screenshot extraction (monthly 汇总 article photos) ───────────────
# The 智慧光伏 monthly article is one tall table covering all provinces. Users
# capture it as a full-page screenshot (~400×30000 px). Vision needs readable
# text, so the scroll is sliced into overlapping chunks; each chunk yields the
# rows it can read fully; chunks are merged by province.

_SCROLL_MONTH_PROMPT = (
    "这是一张微信公众号文章的顶部截图，内容是某年某月中国各省电网代理购电价格汇总。"
    "请判断这篇文章公布的是哪一年哪一月的代理购电价格（标题或开头会写，如「2026年6月」）。"
    "只输出 JSON：{\"year\": 2026, \"month\": 6}。看不出来就输出 {}。"
)

_SCROLL_CHUNK_PROMPT = (
    "这张图片片段来自一篇汇总文章，包含若干省份的代理购电价格公告"
    "（每个省一个【独立的小表格】，表格上方标有省名标题；也有省份只在正文备注里"
    "写「系统运行费用折合度电水平X元/千瓦时」）。\n"
    "对每个【完整可见】的省份表格或备注，提取：\n"
    "1. 省份短名（山东、蒙西、蒙东、冀北、河北南网、深圳…；北京/上海/天津/重庆不带「市」字）\n"
    "2. 该省「系统运行费用折合度电水平」的数值（也可能写作「系统运行费用折价」"
    "「系统运行费折价」「系统运行费」）。\n"
    "⚠️ 绝对不要把代理购电价格、上网环节线损、输配电价、政府性基金等其他行/列的数字"
    "当成系统运行费。找不到该行就不要输出这个省。\n"
    "单位如为分/千瓦时，除以100换算为元/千瓦时。\n"
    "只输出 JSON 数组：[{\"province\": \"山东\", \"fee\": 0.0357}]。"
    "被截断、模糊、看不清的省份【不要】输出。没有可完整读取的省份就输出 []。"
)

_SCROLL_CHUNK_PX = 1000      # ≈ one province block per chunk
_SCROLL_OVERLAP_PX = 200
_SCROLL_UPSCALE = 3          # 394px-wide scrolls need upscaling for legibility


def _vision_json(image_bytes: bytes, prompt: str, api_key: str,
                 max_tokens: int = 2048):
    """Vision call expecting a JSON payload; returns parsed object or None."""
    import json

    from shared.anthropic_client import make_client
    image_bytes = _shrink_image(image_bytes)
    media_type = _sniff_media_type(image_bytes)
    client = make_client(api_key)
    resp = client.messages.create(
        model="claude-sonnet-4-6",
        max_tokens=max_tokens,
        messages=[{
            "role": "user",
            "content": [
                {"type": "image",
                 "source": {"type": "base64", "media_type": media_type,
                            "data": base64.standard_b64encode(image_bytes).decode()}},
                {"type": "text", "text": prompt},
            ],
        }],
    )
    text = resp.content[0].text.strip()
    m = re.search(r"(\[.*\]|\{.*\})", text, re.DOTALL)
    if not m:
        return None
    try:
        return json.loads(m.group(0))
    except json.JSONDecodeError:
        return None


def extract_from_excel_bytes(xlsx_bytes: bytes) -> dict[int, float]:
    """Parse a provincial 代理购电价格汇总及分析 workbook → {month: fee_yuan_kwh}.

    Layout (all 33 provincial files share it): a header row with 1月份…12月份
    columns; a fee row whose 明细 cell is 系统运行费用折价 / 折合度电水平 /
    折价度电水平. Month columns are located from the header, not hardcoded.
    #REF! and empty cells are skipped. 广东/深圳 publish in 分/千瓦时 → the
    >1 guard converts.
    """
    import io

    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = [list(r) for r in ws.iter_rows(min_row=1, max_row=120, values_only=True)]
    wb.close()

    month_cols: dict[int, int] = {}  # month -> column index
    for r in rows:
        for j, c in enumerate(r):
            if isinstance(c, str):
                m = re.fullmatch(r"(\d{1,2})月(?:份)?", c.strip())
                if m:
                    month_cols[int(m.group(1))] = j
        if len(month_cols) >= 12:
            break
    if not month_cols:
        return {}

    label_re = re.compile(r"^\s*系统运行费[用]?(折价度电水平|折合度电水平|折价)\s*$")
    for r in rows:
        if not any(isinstance(c, str) and label_re.match(c) for c in r):
            continue
        out: dict[int, float] = {}
        for month, j in month_cols.items():
            if j >= len(r):
                continue
            v = r[j]
            if isinstance(v, (int, float)):
                out[month] = _normalize_fee(float(v))
        return out
    return {}


def identify_scroll_month(image_bytes: bytes, api_key: str) -> Optional[str]:
    """Read the article month off the top of a scroll screenshot → 'YYYY-MM-01'."""
    import io

    from PIL import Image
    img = Image.open(io.BytesIO(image_bytes))
    top = img.crop((0, 0, img.width, min(1600, img.height)))
    buf = io.BytesIO()
    top.convert("RGB").save(buf, format="JPEG", quality=88)
    obj = _vision_json(buf.getvalue(), _SCROLL_MONTH_PROMPT, api_key, max_tokens=64)
    if not isinstance(obj, dict):
        return None
    try:
        y, mo = int(obj["year"]), int(obj["month"])
        assert 2020 <= y <= 2030 and 1 <= mo <= 12
        return f"{y}-{mo:02d}-01"
    except (KeyError, TypeError, ValueError, AssertionError):
        return None


def slice_scroll(image_bytes: bytes, *, chunk_px: int = _SCROLL_CHUNK_PX,
                 overlap_px: int = _SCROLL_OVERLAP_PX,
                 upscale: int = _SCROLL_UPSCALE) -> list[bytes]:
    """Slice a tall scroll screenshot into overlapping chunks.

    Scroll captures from the phone are ~394px wide; at that width the API
    downscales chunks and 6-decimal table cells smear. Each chunk is therefore
    upscaled `upscale`× and saved as PNG (no JPEG artifacts on digits)."""
    import io

    from PIL import Image
    img = Image.open(io.BytesIO(image_bytes)).convert("RGB")
    w, h = img.size
    step = chunk_px - overlap_px
    chunks = []
    y = 0
    while y < h:
        box = img.crop((0, y, w, min(y + chunk_px, h)))
        if upscale != 1:
            box = box.resize((box.width * upscale, box.height * upscale),
                             Image.LANCZOS)
        buf = io.BytesIO()
        box.save(buf, format="PNG")
        chunks.append(buf.getvalue())
        if y + chunk_px >= h:
            break
        y += step
    return chunks


def extract_scroll_rows(image_bytes: bytes, api_key: str) -> dict[str, float]:
    """Slice a scroll and merge per-chunk {province, fee} rows.

    A province reported by two chunks with different fees is DROPPED (boundary
    misread risk) and logged for manual review — wrong data is worse than a gap.
    """
    merged: dict[str, float] = {}
    conflicts: set[str] = set()
    for i, chunk in enumerate(slice_scroll(image_bytes)):
        rows = _vision_json(chunk, _SCROLL_CHUNK_PROMPT, api_key)
        if not isinstance(rows, list):
            logger.warning("scroll chunk %d: unparseable response", i)
            continue
        for r in rows:
            if not isinstance(r, dict):
                continue
            province = _canon_province_name(str(r.get("province") or ""))
            try:
                fee = _normalize_fee(float(r["fee"]))
            except (KeyError, TypeError, ValueError):
                continue
            if not province:
                continue
            if province in conflicts:
                continue
            prev = merged.get(province)
            if prev is None:
                merged[province] = fee
            elif abs(prev - fee) > 1e-9:
                logger.warning("scroll conflict %s: %.6f vs %.6f — dropped",
                               province, prev, fee)
                conflicts.add(province)
                del merged[province]
    return merged


def _vision_extract_province_fee(image_bytes: bytes, media_type: str,
                                 api_key: str) -> tuple[Optional[str], Optional[float]]:
    """One vision call → (province_short_name, fee) or (None, None)."""
    import json

    from shared.anthropic_client import make_client
    image_bytes = _shrink_image(image_bytes)
    media_type = _sniff_media_type(image_bytes)
    client = make_client(api_key)
    resp = client.messages.create(
        model="claude-sonnet-4-6",
        max_tokens=128,
        messages=[{
            "role": "user",
            "content": [
                {"type": "image",
                 "source": {"type": "base64", "media_type": media_type,
                            "data": base64.standard_b64encode(image_bytes).decode()}},
                {"type": "text", "text": _PROMPT_PROVINCE_FEE},
            ],
        }],
    )
    text = resp.content[0].text.strip()
    m = re.search(r"\{.*\}", text, re.DOTALL)
    if not m:
        return None, None
    try:
        obj = json.loads(m.group(0))
    except json.JSONDecodeError:
        return None, None
    province = _canon_province_name(str(obj.get("province") or ""))
    fee = obj.get("fee")
    try:
        fee = _normalize_fee(float(fee)) if fee is not None else None
    except (TypeError, ValueError):
        fee = None
    return province, fee


def process_article_images(art: dict, *, pg_url: str, api_key: str) -> int:
    """If `art` is a monthly 代理购电价格 announcement, extract per-province
    系统运行费 from its chart images and upsert. Returns rows written (0 when
    the article is not a price announcement — the common case).
    """
    import psycopg2
    import requests

    title = art.get("title") or ""
    body = art.get("body") or ""
    if not body or not _ARTICLE_GATE_RE.search(title + "\n" + body[:2000]):
        return 0

    m = _TITLE_MONTH_RE.search(title)
    if not m:
        logger.info("sysopfee: %r gated on keywords but no 年月 in title — skip", title[:60])
        return 0
    year_month = f"{m.group(1)}-{int(m.group(2)):02d}-01"

    img_urls = [u for u in _IMG_MARKER_RE.findall(body) if _QPIC_HOST_RE.match(u)]
    if not img_urls:
        return 0
    img_urls = img_urls[:_MAX_IMAGES_PER_ARTICLE]

    session = requests.Session()
    session.trust_env = False  # qpic CDN is direct-reachable; ignore proxy env
    session.headers["User-Agent"] = "Mozilla/5.0"

    source_url = art.get("url", "")[:200]
    written = 0
    for url in img_urls:
        try:
            r = session.get(url, timeout=20)
            r.raise_for_status()
            data = r.content
            if len(data) < _MIN_IMAGE_BYTES:
                continue
            media_type = "image/png" if data[:4] == b"\x89PNG" else "image/jpeg"
            province, fee = _vision_extract_province_fee(data, media_type, api_key)
            if not province or fee is None:
                continue
            with psycopg2.connect(pg_url) as conn, conn.cursor() as cur:
                upsert_sysopfee(cur, province, year_month, fee,
                                source_file=source_url or "wechat-article")
            written += 1
            logger.info("sysopfee: %s %s = %.6f ← %s", province, year_month, fee, url[:60])
        except Exception as e:  # noqa: BLE001 — one bad image must not kill the run
            logger.warning("sysopfee: image failed (%s): %s", url[:60], e)
    if written:
        logger.info("sysopfee: %d provinces upserted for %s from %r",
                    written, year_month, title[:60])
    return written

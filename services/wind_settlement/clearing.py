# -*- coding: utf-8 -*-
"""电力交易中心日清算 (daily clearing) parser — the exchange's own per-15-min
settlement record for 悦盛昌渠 (新_悦盛昌渠#1期).

Columns per interval: 计量电量 (metered actual), 电能电费 (energy settlement
amount = 计量×RT + Σ合约×(合约价−区域ref)), 省内实时节点电价, 中长期合约电量/
电价 (aggregate contract curve), 省间日前/日内出清, 曲线合理度取小值/取均值
(= min/mean of 合约 vs 计量 per interval).

This replaces proxy-based replication: metered volumes match book 6 bills to
the MWh, and the interval fee identity lets CfD be derived exactly.
"""
from __future__ import annotations

import datetime as _dt
import warnings
from pathlib import Path

import pandas as pd

COLUMNS = [
    "date", "time", "metered_mwh", "energy_price", "energy_fee",
    "rt_cleared_mw", "rt_nodal_price", "contract_mwh", "contract_price",
    "ic_da_mw", "ic_da_price", "ic_id_mw", "ic_id_price",
    "curve_min", "curve_mean",
]
_NUMERIC = COLUMNS[2:]


def _to_td(v) -> pd.Timedelta:
    if isinstance(v, _dt.datetime):
        return pd.Timedelta(hours=v.hour, minutes=v.minute, seconds=v.second)
    if isinstance(v, _dt.time):
        return pd.Timedelta(hours=v.hour, minutes=v.minute, seconds=v.second)
    parts = str(v).split(":")
    return pd.Timedelta(hours=int(parts[0]), minutes=int(parts[1]),
                        seconds=int(float(parts[2])) if len(parts) > 2 else 0)


def parse_clearing_file(path: str | Path) -> pd.DataFrame:
    """Parse the 日清算 workbook into the COLUMNS schema + combined datetime."""
    import openpyxl

    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        wb = openpyxl.load_workbook(path, read_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = [r for i, r in enumerate(ws.iter_rows(values_only=True)) if i > 0]

    df = pd.DataFrame(rows, columns=COLUMNS)
    for c in _NUMERIC:
        df[c] = pd.to_numeric(df[c], errors="coerce")
    df["datetime"] = pd.to_datetime(df["date"]) + df["time"].map(_to_td)
    df = df.sort_values("datetime").reset_index(drop=True)
    df["source_file"] = Path(path).name
    return df


def monthly_metered(clearing: pd.DataFrame) -> pd.DataFrame:
    """Monthly aggregates: metered volume, energy fee, contract volume, Σmin."""
    m = clearing.copy()
    m["month"] = m["datetime"].dt.strftime("%Y-%m")
    return m.groupby("month").agg(
        metered_mwh=("metered_mwh", "sum"),
        energy_fee=("energy_fee", "sum"),
        contract_mwh=("contract_mwh", "sum"),
        curve_min=("curve_min", "sum"),
    )

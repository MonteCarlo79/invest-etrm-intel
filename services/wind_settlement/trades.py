# -*- coding: utf-8 -*-
"""零碳46 (悦盛昌渠) medium/long-term trade confirmations parser.

Parses the monthly 成交单总表 workbooks into a normalised DataFrame:

- 省内 files (3.省内成交单分月总表/): 15-column schema, includes per-品种
  汇总 subtotal rows (net of 置换 swaps) and a 全部交易品种 grand-total row.
- 跨省 files (2.跨省成交单分月总表/): 18-column schema, plain rows, no subtotals.

Volumes are monthly *delivery* slices (annual contracts appear in every
monthly file with that month's volume), so rows can feed settlement
replication directly.
"""
from __future__ import annotations

import re
from pathlib import Path

import pandas as pd

COLUMNS = [
    "row_no", "trade_type", "mode", "period", "energy_kind",
    "consumer_unit", "generator_unit", "volume_mwh",
    "energy_price", "energy_fee", "env_value", "all_in_price",
]

_INTRA_IDX = {  # 省内 15-col layout
    "row_no": 0, "trade_type": 1, "mode": 2, "period": 3, "energy_kind": 4,
    "consumer_unit": 5, "generator_unit": 6, "volume_mwh": 10,
    "energy_price": 11, "energy_fee": 12, "env_value": 13, "all_in_price": 14,
}
_CROSS_IDX = {  # 跨省 18-col layout
    "row_no": 0, "trade_type": 1, "mode": 2, "period": 3, "energy_kind": 4,
    "consumer_unit": 6, "generator_unit": 8, "volume_mwh": 13,
    "energy_price": 14, "energy_fee": 17, "env_value": 15, "all_in_price": 16,
}

_MONTH_RE = re.compile(r"(20\d{2})-(\d{2})")


def _num(v) -> float | None:
    if v is None or v == "":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def parse_trades_file(path: str | Path) -> pd.DataFrame:
    """Parse one 成交单总表 workbook into the normalised COLUMNS schema."""
    import openpyxl

    path = Path(path)
    m = _MONTH_RE.search(path.name)
    if not m:
        raise ValueError(f"cannot read delivery month from filename: {path.name}")
    delivery_month = f"{m.group(1)}-{m.group(2)}"
    channel = "跨省" if "跨省" in str(path.parent) else "省内"
    idx = _CROSS_IDX if channel == "跨省" else _INTRA_IDX

    import warnings
    with warnings.catch_warnings():
        # the trading-platform exports carry no default style sheet
        warnings.simplefilter("ignore", UserWarning)
        wb = openpyxl.load_workbook(path)
    ws = wb[wb.sheetnames[0]]
    rows = []
    for i, raw in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue  # header
        rec = {name: raw[pos] if pos < len(raw) else None for name, pos in idx.items()}
        if rec["trade_type"] is None:
            continue
        for f in ("volume_mwh", "energy_price", "energy_fee", "env_value", "all_in_price"):
            rec[f] = _num(rec[f])
        rows.append(rec)

    df = pd.DataFrame(rows, columns=COLUMNS)
    df["delivery_month"] = delivery_month
    df["channel"] = channel
    df["source_file"] = path.name
    return df


def monthly_position(intra: pd.DataFrame, cross: pd.DataFrame) -> pd.DataFrame:
    """Net monthly contract position: 省内 per-品种 汇总 rows + all 跨省 rows.

    The 汇总 rows are the trading platform's own net-of-swap figures, so they
    are authoritative for the month's delivery position.
    """
    parts = []
    if intra is not None and not intra.empty:
        summ = intra[
            (intra["energy_kind"] == "汇总")
            & (~intra["trade_type"].str.startswith("全部交易品种"))
        ]
        parts.append(summ)
    if cross is not None and not cross.empty:
        parts.append(cross)

    out = []
    for part in parts:
        for _, grp in part.groupby(["channel", "trade_type"]):
            vol = grp["volume_mwh"].sum()
            if vol == 0:
                continue
            wavg = lambda col: float((grp[col] * grp["volume_mwh"]).sum() / vol) if grp[col].notna().any() else None  # noqa: E731
            out.append({
                "channel": grp["channel"].iloc[0],
                "trade_type": grp["trade_type"].iloc[0],
                "mode": grp["mode"].iloc[0],
                "period": grp["period"].iloc[0],
                "volume_mwh": float(vol),
                "energy_price": wavg("energy_price"),
                "env_value": wavg("env_value"),
            })
    return pd.DataFrame(out)

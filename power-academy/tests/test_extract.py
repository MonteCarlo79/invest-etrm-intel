import json

import docx
import pymupdf
from pptx import Presentation

from academy.extract import FOLDER_CAP, extract_all, extract_source
from academy.models import SourceEntry


def _entry(id_, path, type_="pdf", cls="library"):
    return SourceEntry(id=id_, path=str(path), title=id_, type=type_,
                       source_class=cls, cleared=cls == "practice",
                       license_risk="high" if cls == "library" else "low")


def _pdf(path, text):
    d = pymupdf.open()
    page = d.new_page()
    if text:
        page.insert_textbox(pymupdf.Rect(50, 50, 550, 780), text)
    d.save(path)


def test_pdf_ok_and_blank_pdf_is_no_text(tmp_path):
    _pdf(tmp_path / "a.pdf", "spark spread option " * 30)
    _pdf(tmp_path / "blank.pdf", "")
    cache = tmp_path / "cache"
    ok = extract_source(_entry("a", tmp_path / "a.pdf"), cache)
    blank = extract_source(_entry("blank", tmp_path / "blank.pdf"), cache)
    assert ok["status"] == "ok" and (cache / "a.txt").exists()
    assert blank["status"] == "no_text" and not (cache / "blank.txt").exists()


def test_pptx_docx_txt(tmp_path):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "tolling agreement " * 20
    prs.save(tmp_path / "d.pptx")
    dd = docx.Document()
    dd.add_paragraph("storage valuation " * 20)
    dd.save(tmp_path / "d.docx")
    (tmp_path / "d.txt").write_text("swing option " * 40, encoding="utf-8")
    cache = tmp_path / "cache"
    for stem, t in [("d", "pptx"), ("d", "docx"), ("d", "txt")]:
        e = _entry(f"id_{t}", tmp_path / f"{stem}.{t}", t)
        assert extract_source(e, cache)["status"] == "ok"


def test_legacy_ppt_unsupported_and_corrupt_failed(tmp_path):
    (tmp_path / "old.ppt").write_bytes(b"x")
    (tmp_path / "bad.pdf").write_bytes(b"not a pdf")
    cache = tmp_path / "cache"
    assert extract_source(_entry("old", tmp_path / "old.ppt", "ppt"), cache)["status"] == "unsupported"
    assert extract_source(_entry("bad", tmp_path / "bad.pdf"), cache)["status"] == "failed"


def test_practice_folder_docs_models_only_and_cap(tmp_path):
    f = tmp_path / "deal"
    f.mkdir()
    (f / "notes.txt").write_text("hedge " * 100, encoding="utf-8")
    (f / "old.ppt").write_bytes(b"x")
    cache = tmp_path / "cache"
    r = extract_source(_entry("deal", f, "folder", "practice"), cache)
    assert r["status"] == "ok"
    body = (cache / "deal.txt").read_text(encoding="utf-8")
    assert "[skipped: old.ppt" in body

    only_models = tmp_path / "models"
    only_models.mkdir()
    (only_models / "m.xlsx").write_bytes(b"x")
    assert extract_source(_entry("models", only_models, "folder", "practice"), cache)["status"] == "no_text"

    big = tmp_path / "big"
    big.mkdir()
    for i in range(3):
        (big / f"{i}.txt").write_text("x" * 150_000, encoding="utf-8")
    assert extract_source(_entry("big", big, "folder", "practice"), cache)["chars"] <= FOLDER_CAP + 500


def test_extract_all_counts_cached_and_report(tmp_path):
    _pdf(tmp_path / "a.pdf", "forward curve " * 40)
    (tmp_path / "old.ppt").write_bytes(b"x")
    es = [_entry("a", tmp_path / "a.pdf"), _entry("old", tmp_path / "old.ppt", "ppt")]
    cache = tmp_path / "cache"
    r1 = extract_all(es, cache, throttle_s=0)
    assert r1["counts"] == {"ok": 1, "unsupported": 1}
    r2 = extract_all(es, cache, throttle_s=0)
    assert r2["counts"] == {"cached": 1, "unsupported": 1}
    assert json.loads((cache / "_report.json").read_text())["counts"] == r2["counts"]

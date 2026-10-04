import hashlib
import json
import os
import shutil
import subprocess
import time
from pathlib import Path

import docx
import pymupdf
from pptx import Presentation

from .registry import DOC_EXT, MODEL_EXT, walk_files

MIN_CHARS = 200
MIN_PAGE_CHARS = 20  # below this a PDF page is treated as scanned (OCR candidate)
FOLDER_CAP = 200_000


LEGACY = {".ppt": ".pptx", ".doc": ".docx"}
_SOFFICE_PATHS = [Path.home() / "Applications/LibreOffice.app/Contents/MacOS/soffice",
                  Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")]


class Unsupported(Exception):
    pass


def find_soffice():
    found = os.environ.get("SOFFICE") or shutil.which("soffice")
    if found:
        return found
    return next((str(p) for p in _SOFFICE_PATHS if p.exists()), None)


def convert_legacy(path: Path, convert_dir: Path) -> Path:
    """Convert .ppt/.doc to .pptx/.docx inside convert_dir (never beside the source)."""
    path = Path(path)
    target = LEGACY[path.suffix.lower()]
    out_dir = Path(convert_dir) / hashlib.sha1(str(path).encode("utf-8")).hexdigest()[:10]
    out = out_dir / (path.stem + target)
    if out.exists():
        return out
    soffice = find_soffice()
    if soffice is None:
        raise Unsupported(path.suffix.lower())
    out_dir.mkdir(parents=True, exist_ok=True)
    r = subprocess.run(
        [soffice, "-env:UserInstallation=file:///tmp/pa_lo_profile", "--headless",
         "--convert-to", target[1:], "--outdir", str(out_dir), str(path)],
        capture_output=True, timeout=180)
    if r.returncode != 0 or not out.exists():
        raise RuntimeError(f"LibreOffice conversion failed (exit {r.returncode})")
    return out


def _page_text(page, ocr: bool) -> str:
    text = page.get_text()
    if ocr and len(text.strip()) < MIN_PAGE_CHARS:
        from .ocr import ocr_page
        return ocr_page(page)
    return text


def extract_file(path: Path, convert_dir=None, ocr: bool = False) -> str:
    ext = Path(path).suffix.lower()
    if ext in LEGACY:
        if convert_dir is None:
            raise Unsupported(ext)
        path = convert_legacy(path, convert_dir)
        ext = Path(path).suffix.lower()
    if ext == ".pdf":
        with pymupdf.open(path) as d:
            return "\n".join(_page_text(page, ocr) for page in d)
    if ext == ".pptx":
        out = []
        for i, slide in enumerate(Presentation(str(path)).slides, 1):
            out.append(f"[slide {i}]")
            out.extend(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)
        return "\n".join(out)
    if ext == ".docx":
        return "\n".join(p.text for p in docx.Document(str(path)).paragraphs)
    if ext == ".txt":
        return Path(path).read_text(encoding="utf-8", errors="replace")
    raise Unsupported(ext)


def _extract_folder(folder: Path, convert_dir=None, ocr: bool = False) -> str:
    parts, total, models = [], 0, []
    for p in walk_files(folder):
        ext = p.suffix.lower()
        if ext in MODEL_EXT:
            models.append(p.name)
            continue
        if ext not in DOC_EXT:
            continue
        try:
            body = extract_file(p, convert_dir, ocr)
        except Unsupported:
            parts.append(f"[skipped: {p.name} (unsupported {ext})]")
            continue
        except Exception as e:  # corrupt file: note and continue
            parts.append(f"[skipped: {p.name} ({type(e).__name__})]")
            continue
        chunk = f"=== {p.relative_to(folder)} ===\n{body}"[:FOLDER_CAP - total]
        parts.append(chunk)
        total += len(chunk)
        if total >= FOLDER_CAP:
            parts.append("[truncated: folder cap reached]")
            break
    if models:
        parts.append("[models] " + ", ".join(models))
    return "\n".join(parts)


def extract_source(entry, cache_dir: Path, convert_dir=None, ocr: bool = False) -> dict:
    cache_dir = Path(cache_dir)
    out = cache_dir / f"{entry.id}.txt"
    res = {"id": entry.id, "status": "ok", "chars": 0, "error": None}
    if out.exists():
        res["status"] = "cached"
        res["chars"] = len(out.read_text(encoding="utf-8"))
        return res
    try:
        text = _extract_folder(Path(entry.path), convert_dir, ocr) if entry.type == "folder" \
            else extract_file(Path(entry.path), convert_dir, ocr)
    except Unsupported as e:
        return {**res, "status": "unsupported", "error": f"unsupported {e}"}
    except Exception as e:
        return {**res, "status": "failed", "error": f"{type(e).__name__}: {e}"}
    res["chars"] = len(text.strip())
    if res["chars"] < MIN_CHARS:
        res["status"] = "no_text"
        return res
    cache_dir.mkdir(parents=True, exist_ok=True)
    out.write_text(text, encoding="utf-8")
    return res


def extract_all(entries, cache_dir: Path, throttle_s: float = 0.2, convert_dir=None,
                ocr: bool = False) -> dict:
    results, counts = [], {}
    for e in entries:
        r = extract_source(e, cache_dir, convert_dir, ocr)
        results.append(r)
        counts[r["status"]] = counts.get(r["status"], 0) + 1
        if r["status"] != "cached":
            time.sleep(throttle_s)
    report = {"counts": counts, "results": results}
    Path(cache_dir).mkdir(parents=True, exist_ok=True)
    (Path(cache_dir) / "_report.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=1), encoding="utf-8")
    return report
